from pathlib import Path
from xml.sax.saxutils import escape

from openpyxl import load_workbook
from PIL import Image, ImageDraw, ImageFilter, ImageFont


ROOT = Path("/Users/zhonghao/Projects/my-tampermonkey-scripts")
SOURCE = Path("/Users/zhonghao/Downloads/【人天评估】Ai 学习助手项目-功能确认单V4.xlsx")
OUT = ROOT / "output"

PNG = OUT / "ai-learning-assistant-architecture.png"
SVG = OUT / "ai-learning-assistant-architecture.svg"
MD = OUT / "ai-learning-assistant-architecture-philosophy.md"

W, H = 3840, 2160

BG_TOP = (248, 250, 252)
BG_BOTTOM = (239, 244, 248)
INK = (28, 34, 44)
MUTED = (97, 108, 124)
WHITE = (255, 255, 255)

ACCENTS = [
    (0, 151, 167),
    (235, 112, 86),
    (108, 91, 216),
    (47, 148, 107),
    (222, 151, 46),
]

LIGHTS = [
    (228, 248, 250),
    (255, 239, 235),
    (240, 237, 255),
    (232, 248, 240),
    (255, 246, 229),
]


def font(size, bold=False):
    candidates = [
        "/System/Library/Fonts/STHeiti Medium.ttc" if bold else "/System/Library/Fonts/STHeiti Light.ttc",
        "/System/Library/Fonts/Hiragino Sans GB.ttc",
        "/Library/Fonts/Arial Unicode.ttf",
        "/System/Library/Fonts/Helvetica.ttc",
    ]
    for path in candidates:
        try:
            return ImageFont.truetype(path, size=size)
        except OSError:
            continue
    return ImageFont.load_default()


F_TITLE = font(82, True)
F_SUBTITLE = font(34)
F_META = font(24)
F_PANEL = font(34, True)
F_PANEL_SMALL = font(22)
F_CARD_TITLE = font(29, True)
F_CARD_BODY = font(23)
F_FOOT = font(31, True)
F_FOOT_SMALL = font(24)


PANELS = [
    {
        "title": "用户端 App",
        "subtitle": "页面体验与核心入口",
        "icon": "app",
        "cards": [
            ("登录与套餐激活", "手机号验证码、一键登录、首次登录弹窗、额度初始化"),
            ("首页工作台", "套餐额度看板、模型入口、智能体快捷入口、新手指引"),
            ("模型广场", "分类筛选、参数对比、模型状态、立即创建会话"),
            ("智能体中心", "岗位/场景分类、全局搜索、Prompt 模板与表单联动"),
            ("AI 聊天与个人中心", "会话历史、SSE 打字机、用量趋势、API Key 入口"),
        ],
    },
    {
        "title": "接入与安全",
        "subtitle": "统一入口、身份与风控",
        "icon": "shield",
        "cards": [
            ("HTTPS / SSL", "全站加密，生产环境安全证书与访问保护"),
            ("统一认证", "手机号校验、验证码倒计时、防刷与每日上限"),
            ("API Key 管理", "密钥创建、启停、脱敏展示，兼容开发者接入"),
            ("限额与熔断", "请求前额度检查，耗尽置灰，网关返回 402 拦截"),
            ("安全存储", "前端摘要配合后端 Bcrypt，敏感信息脱敏展示"),
        ],
    },
    {
        "title": "AI 能力网关",
        "subtitle": "模型路由、流式响应与计量",
        "icon": "gateway",
        "cards": [
            ("模型智能路由", "基于百炼模型列表与 model_id 分发 DeepSeek、Qwen 等模型"),
            ("SSE 流式对话", "逐帧输出 Markdown，支持代码复制、公式表格与停止生成"),
            ("Token 计量扣减", "在最终 usage 帧汇总 Prompt + Completion，准实时扣减额度"),
            ("OpenAI 协议兼容", "把平台额度桥接到 Dify、Cursor 等第三方工具"),
        ],
    },
    {
        "title": "业务服务层",
        "subtitle": "可配置、可运营的中台能力",
        "icon": "service",
        "cards": [
            ("账号套餐服务", "套餐激活、总额度、剩余额度、到期时间与倒计时"),
            ("模型目录服务", "模型分类、标签、上下文窗口、消耗系数与维护状态"),
            ("智能体模板服务", "System Prompt、用户表单、快捷入口与场景配置"),
            ("会话消息服务", "新建会话、历史分组、上下文加载、重命名与删除"),
            ("用量统计服务", "最近 7/30 天 Token 趋势、调用次数、首页看板刷新"),
        ],
    },
    {
        "title": "数据与外部能力",
        "subtitle": "持久化、模型平台与交付保障",
        "icon": "data",
        "cards": [
            ("业务数据库", "用户、套餐、API Key、消息、用量流水与模型配置"),
            ("缓存与日志", "会话上下文缓存、请求日志、计量记录与问题追踪"),
            ("阿里百炼平台", "模型列表 API、DeepSeek、通义千问、多模态能力接入"),
            ("运维交付", "服务器部署、压力测试、安全测试、12 个月维护服务"),
        ],
    },
]


def lerp(a, b, t):
    return int(a + (b - a) * t)


def text_width(draw, text, fnt):
    box = draw.textbbox((0, 0), text, font=fnt)
    return box[2] - box[0]


def line_height(fnt, extra=8):
    box = fnt.getbbox("天地Ag")
    return box[3] - box[1] + extra


def wrap_text(draw, text, fnt, max_width):
    lines = []
    for part in str(text).split("\n"):
        current = ""
        for ch in part:
            trial = current + ch
            if not current or text_width(draw, trial, fnt) <= max_width:
                current = trial
            else:
                lines.append(current)
                current = ch
        if current:
            lines.append(current)
    return lines or [""]


def truncate_to_width(draw, text, fnt, max_width):
    if text_width(draw, text, fnt) <= max_width:
        return text
    suffix = "..."
    text = text.rstrip()
    while text and text_width(draw, text + suffix, fnt) > max_width:
        text = text[:-1]
    return text + suffix


def draw_wrapped(draw, xy, text, fnt, fill, max_width, max_lines=2, gap=6):
    x, y = xy
    lines = wrap_text(draw, text, fnt, max_width)
    if len(lines) > max_lines:
        lines = lines[:max_lines]
        lines[-1] = truncate_to_width(draw, lines[-1], fnt, max_width)
    lh = line_height(fnt, gap)
    for i, line in enumerate(lines):
        draw.text((x, y + i * lh), line, font=fnt, fill=fill)
    return y + len(lines) * lh


def rounded_with_shadow(img, box, radius, fill, outline=None, width=2, shadow=True):
    if shadow:
        layer = Image.new("RGBA", img.size, (0, 0, 0, 0))
        sd = ImageDraw.Draw(layer)
        x1, y1, x2, y2 = box
        sd.rounded_rectangle((x1 + 8, y1 + 12, x2 + 8, y2 + 12), radius=radius, fill=(28, 39, 54, 24))
        layer = layer.filter(ImageFilter.GaussianBlur(18))
        img.alpha_composite(layer)
    draw = ImageDraw.Draw(img)
    draw.rounded_rectangle(box, radius=radius, fill=fill, outline=outline, width=width)


def draw_arrow(draw, start, end, color, width=5):
    x1, y1 = start
    x2, y2 = end
    draw.line((x1, y1, x2, y2), fill=color, width=width)
    head = 16
    draw.polygon([(x2, y2), (x2 - head, y2 - head // 2), (x2 - head, y2 + head // 2)], fill=color)


def draw_icon(draw, kind, cx, cy, color):
    r = 34
    draw.ellipse((cx - r, cy - r, cx + r, cy + r), fill=color)
    c = WHITE
    lw = 5
    if kind == "app":
        draw.rounded_rectangle((cx - 16, cy - 23, cx + 16, cy + 23), radius=7, outline=c, width=lw)
        draw.line((cx - 8, cy + 14, cx + 8, cy + 14), fill=c, width=lw)
    elif kind == "shield":
        draw.polygon(
            [(cx, cy - 24), (cx + 21, cy - 14), (cx + 16, cy + 17), (cx, cy + 28), (cx - 16, cy + 17), (cx - 21, cy - 14)],
            outline=c,
            fill=None,
        )
        draw.line((cx - 10, cy, cx - 1, cy + 9, cx + 13, cy - 10), fill=c, width=lw)
    elif kind == "gateway":
        draw.line((cx - 22, cy - 14, cx + 22, cy - 14), fill=c, width=lw)
        draw.line((cx - 22, cy + 14, cx + 22, cy + 14), fill=c, width=lw)
        draw.line((cx - 11, cy - 24, cx + 11, cy + 24), fill=c, width=lw)
    elif kind == "service":
        for dx, dy in [(-14, -14), (14, -14), (-14, 14), (14, 14)]:
            draw.rounded_rectangle((cx + dx - 10, cy + dy - 10, cx + dx + 10, cy + dy + 10), radius=4, outline=c, width=lw)
        draw.line((cx - 4, cy - 14, cx + 4, cy - 14), fill=c, width=lw)
        draw.line((cx - 4, cy + 14, cx + 4, cy + 14), fill=c, width=lw)
    else:
        draw.ellipse((cx - 22, cy - 14, cx + 22, cy + 14), outline=c, width=lw)
        draw.line((cx - 22, cy - 14, cx - 22, cy + 16), fill=c, width=lw)
        draw.line((cx + 22, cy - 14, cx + 22, cy + 16), fill=c, width=lw)
        draw.arc((cx - 22, cy + 2, cx + 22, cy + 30), 0, 180, fill=c, width=lw)


def counts_from_workbook():
    try:
        ws = load_workbook(SOURCE, data_only=True).active
    except Exception:
        return "功能确认单 V4 · 用户端功能 24 项 · 服务交付 14 项"
    nav = None
    user_features = 0
    navs = set()
    services = 0
    for idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
        _, b, c, d = (list(row) + [None, None, None, None])[:4]
        if b:
            nav = b
        if 7 <= idx <= 30 and c:
            user_features += 1
            if nav:
                navs.add(nav)
        if 33 <= idx <= 46 and d:
            services += 1
    return f"功能确认单 V4 · 用户端功能 {user_features} 项 · 导航 {len(navs)} 组 · 服务交付 {services} 项"


def draw_png():
    img = Image.new("RGBA", (W, H), BG_TOP + (255,))
    px = img.load()
    for y in range(H):
        t = y / (H - 1)
        color = tuple(lerp(BG_TOP[i], BG_BOTTOM[i], t) for i in range(3)) + (255,)
        for x in range(W):
            px[x, y] = color

    draw = ImageDraw.Draw(img)
    for x in range(120, W - 120, 96):
        draw.line((x, 300, x, H - 260), fill=(232, 238, 244), width=1)
    for y in range(330, H - 260, 96):
        draw.line((120, y, W - 120, y), fill=(232, 238, 244), width=1)
    draw.line((150, 232, W - 150, 232), fill=(218, 226, 235), width=2)

    meta = counts_from_workbook()
    draw.text((150, 100), "AI 学习助手系统架构图", font=F_TITLE, fill=INK)
    draw.text((154, 205), "用户体验入口 · AI 网关计量 · 业务中台服务 · 模型与数据基础设施", font=F_SUBTITLE, fill=MUTED)
    m_w = text_width(draw, meta, F_META)
    draw.rounded_rectangle((W - 150 - m_w - 48, 100, W - 150, 150), radius=25, fill=(255, 255, 255, 210), outline=(220, 228, 238), width=2)
    draw.text((W - 150 - m_w - 24, 112), meta, font=F_META, fill=MUTED)

    panel_w = 680
    gap = 38
    left = (W - panel_w * 5 - gap * 4) // 2
    top = 345
    panel_h = 1395
    panel_boxes = []

    for i, panel in enumerate(PANELS):
        x = left + i * (panel_w + gap)
        y = top
        panel_boxes.append((x, y, x + panel_w, y + panel_h))
        rounded_with_shadow(img, (x, y, x + panel_w, y + panel_h), 34, WHITE, outline=(218, 226, 236), width=2, shadow=True)
        draw = ImageDraw.Draw(img)
        accent = ACCENTS[i]
        light = LIGHTS[i]
        draw.rounded_rectangle((x + 24, y + 24, x + panel_w - 24, y + 126), radius=24, fill=light)
        draw.rounded_rectangle((x + 24, y + 24, x + 36, y + 126), radius=6, fill=accent)
        draw_icon(draw, panel["icon"], x + 82, y + 75, accent)
        draw.text((x + 138, y + 44), f"{i + 1:02d}  {panel['title']}", font=F_PANEL, fill=INK)
        draw.text((x + 140, y + 91), panel["subtitle"], font=F_PANEL_SMALL, fill=MUTED)

        card_gap = 18
        card_count = len(panel["cards"])
        cards_top = y + 156
        cards_bottom = y + panel_h - 36
        card_h = int((cards_bottom - cards_top - card_gap * (card_count - 1)) / card_count)
        for j, (title, body) in enumerate(panel["cards"]):
            cy = cards_top + j * (card_h + card_gap)
            cx = x + 34
            cw = panel_w - 68
            draw.rounded_rectangle((cx, cy, cx + cw, cy + card_h), radius=22, fill=(252, 253, 255), outline=(224, 231, 239), width=2)
            draw.rounded_rectangle((cx + 20, cy + 25, cx + 66, cy + 71), radius=15, fill=light)
            draw.text((cx + 32, cy + 31), str(j + 1), font=font(25, True), fill=accent)
            draw.text((cx + 84, cy + 24), title, font=F_CARD_TITLE, fill=INK)
            draw_wrapped(draw, (cx + 84, cy + 70), body, F_CARD_BODY, MUTED, cw - 112, max_lines=3, gap=7)

    for i in range(4):
        x1, y1, x2, y2 = panel_boxes[i]
        nx1, ny1, nx2, ny2 = panel_boxes[i + 1]
        mid_y = top + panel_h // 2
        draw_arrow(draw, (x2 + 6, mid_y), (nx1 - 8, mid_y), (172, 184, 198), width=5)

    flow_x = 220
    flow_y = 1844
    flow_w = W - 440
    flow_h = 176
    rounded_with_shadow(img, (flow_x, flow_y, flow_x + flow_w, flow_y + flow_h), 34, (31, 39, 50), outline=None, width=0, shadow=True)
    draw = ImageDraw.Draw(img)
    draw.text((flow_x + 48, flow_y + 42), "核心业务闭环", font=F_FOOT, fill=WHITE)
    draw.text((flow_x + 48, flow_y + 92), "从套餐激活到模型调用，再到额度扣减和看板刷新，形成可计量、可运营的 AI 服务链路", font=F_FOOT_SMALL, fill=(203, 213, 225))

    steps = ["登录激活", "选模型/智能体", "SSE 对话", "Token 扣减", "额度刷新"]
    sx = flow_x + 760
    pill_w = 410
    pill_gap = 38
    for i, step in enumerate(steps):
        x = sx + i * (pill_w + pill_gap)
        y = flow_y + 55
        draw.rounded_rectangle((x, y, x + pill_w, y + 72), radius=36, fill=(255, 255, 255, 235))
        draw.ellipse((x + 18, y + 17, x + 56, y + 55), fill=ACCENTS[i % len(ACCENTS)])
        draw.text((x + 31, y + 21), str(i + 1), font=font(20, True), fill=WHITE)
        tw = text_width(draw, step, F_CARD_TITLE)
        draw.text((x + 72 + (pill_w - 92 - tw) / 2, y + 18), step, font=F_CARD_TITLE, fill=INK)
        if i < len(steps) - 1:
            draw_arrow(draw, (x + pill_w + 8, y + 36), (x + pill_w + pill_gap - 10, y + 36), (151, 164, 181), width=4)

    draw.text((150, H - 88), "架构提炼自《Ai 学习助手项目-功能确认单 V4》", font=F_META, fill=(116, 128, 144))
    draw.text((W - 150 - text_width(draw, "PPT 立项版 · 16:9 高清图", F_META), H - 88), "PPT 立项版 · 16:9 高清图", font=F_META, fill=(116, 128, 144))
    img.convert("RGB").save(PNG, quality=96)


def approx_wrap(text, max_units):
    lines = []
    current = ""
    width = 0.0
    for ch in text:
        units = 0.56 if ord(ch) < 128 else 1.0
        if current and width + units > max_units:
            lines.append(current)
            current = ch
            width = units
        else:
            current += ch
            width += units
    if current:
        lines.append(current)
    return lines


def svg_text(x, y, lines, size, color, weight=400, line_gap=1.25):
    out = [f'<text x="{x}" y="{y}" font-size="{size}" font-weight="{weight}" fill="{color}">']
    for i, line in enumerate(lines):
        dy = 0 if i == 0 else size * line_gap
        out.append(f'<tspan x="{x}" dy="{dy}">{escape(line)}</tspan>')
    out.append("</text>")
    return "".join(out)


def rgb(c):
    return f"rgb({c[0]},{c[1]},{c[2]})"


def draw_svg():
    panel_w = 680
    gap = 38
    left = (W - panel_w * 5 - gap * 4) // 2
    top = 345
    panel_h = 1395
    parts = [
        f'<svg xmlns="http://www.w3.org/2000/svg" width="{W}" height="{H}" viewBox="0 0 {W} {H}">',
        "<defs>",
        '<linearGradient id="bg" x1="0" y1="0" x2="0" y2="1"><stop offset="0%" stop-color="#F8FAFC"/><stop offset="100%" stop-color="#EFF4F8"/></linearGradient>',
        '<filter id="shadow" x="-10%" y="-10%" width="120%" height="130%"><feDropShadow dx="8" dy="14" stdDeviation="14" flood-color="#1c2736" flood-opacity="0.12"/></filter>',
        "<style>text{font-family:'Hiragino Sans GB','PingFang SC','Microsoft YaHei',Arial,sans-serif;letter-spacing:0}</style>",
        "</defs>",
        '<rect width="3840" height="2160" fill="url(#bg)"/>',
    ]
    for x in range(120, W - 120, 96):
        parts.append(f'<line x1="{x}" y1="300" x2="{x}" y2="1900" stroke="#E8EEF4" stroke-width="1"/>')
    for y in range(330, H - 260, 96):
        parts.append(f'<line x1="120" y1="{y}" x2="3720" y2="{y}" stroke="#E8EEF4" stroke-width="1"/>')
    parts.append('<line x1="150" y1="232" x2="3690" y2="232" stroke="#DAE2EB" stroke-width="2"/>')
    parts.append(svg_text(150, 164, ["AI 学习助手系统架构图"], 82, "#1C222C", 700))
    parts.append(svg_text(154, 242, ["用户体验入口 · AI 网关计量 · 业务中台服务 · 模型与数据基础设施"], 34, "#616C7C", 400))
    meta = counts_from_workbook()
    parts.append('<rect x="2740" y="100" width="950" height="50" rx="25" fill="white" fill-opacity="0.82" stroke="#DCE4EE" stroke-width="2"/>')
    parts.append(svg_text(2770, 134, [meta], 24, "#616C7C"))

    centers = []
    for i, panel in enumerate(PANELS):
        x = left + i * (panel_w + gap)
        y = top
        centers.append((x + panel_w, x, y + panel_h // 2))
        accent = rgb(ACCENTS[i])
        light = rgb(LIGHTS[i])
        parts.append(f'<rect x="{x}" y="{y}" width="{panel_w}" height="{panel_h}" rx="34" fill="white" stroke="#DAE2EC" stroke-width="2" filter="url(#shadow)"/>')
        parts.append(f'<rect x="{x + 24}" y="{y + 24}" width="{panel_w - 48}" height="102" rx="24" fill="{light}"/>')
        parts.append(f'<rect x="{x + 24}" y="{y + 24}" width="12" height="102" rx="6" fill="{accent}"/>')
        parts.append(f'<circle cx="{x + 82}" cy="{y + 75}" r="34" fill="{accent}"/>')
        parts.append(svg_text(x + 138, y + 75, [f"{i + 1:02d}  {panel['title']}"], 34, "#1C222C", 700))
        parts.append(svg_text(x + 140, y + 116, [panel["subtitle"]], 22, "#616C7C"))

        card_count = len(panel["cards"])
        card_gap = 18
        cards_top = y + 156
        cards_bottom = y + panel_h - 36
        card_h = int((cards_bottom - cards_top - card_gap * (card_count - 1)) / card_count)
        for j, (title, body) in enumerate(panel["cards"]):
            cy = cards_top + j * (card_h + card_gap)
            cx = x + 34
            cw = panel_w - 68
            parts.append(f'<rect x="{cx}" y="{cy}" width="{cw}" height="{card_h}" rx="22" fill="#FCFDFF" stroke="#E0E7EF" stroke-width="2"/>')
            parts.append(f'<rect x="{cx + 20}" y="{cy + 25}" width="46" height="46" rx="15" fill="{light}"/>')
            parts.append(svg_text(cx + 33, cy + 58, [str(j + 1)], 25, accent, 700))
            parts.append(svg_text(cx + 84, cy + 56, [title], 29, "#1C222C", 700))
            max_units = (cw - 112) / 25
            lines = approx_wrap(body, max_units)[:3]
            if len(approx_wrap(body, max_units)) > 3:
                lines[-1] = lines[-1][: max(1, len(lines[-1]) - 2)] + "..."
            parts.append(svg_text(cx + 84, cy + 101, lines, 23, "#616C7C", 400, line_gap=1.35))

    for i in range(4):
        x2, _, mid_y = centers[i]
        _, nx, _ = centers[i + 1]
        parts.append(f'<line x1="{x2 + 6}" y1="{mid_y}" x2="{nx - 24}" y2="{mid_y}" stroke="#ACB8C6" stroke-width="5"/>')
        parts.append(f'<polygon points="{nx - 8},{mid_y} {nx - 24},{mid_y - 8} {nx - 24},{mid_y + 8}" fill="#ACB8C6"/>')

    flow_x = 220
    flow_y = 1844
    flow_w = W - 440
    flow_h = 176
    parts.append(f'<rect x="{flow_x}" y="{flow_y}" width="{flow_w}" height="{flow_h}" rx="34" fill="#1F2732" filter="url(#shadow)"/>')
    parts.append(svg_text(flow_x + 48, flow_y + 80, ["核心业务闭环"], 31, "white", 700))
    parts.append(svg_text(flow_x + 48, flow_y + 126, ["从套餐激活到模型调用，再到额度扣减和看板刷新，形成可计量、可运营的 AI 服务链路"], 24, "#CBD5E1"))
    steps = ["登录激活", "选模型/智能体", "SSE 对话", "Token 扣减", "额度刷新"]
    sx = flow_x + 760
    pill_w = 410
    pill_gap = 38
    for i, step in enumerate(steps):
        x = sx + i * (pill_w + pill_gap)
        y = flow_y + 55
        parts.append(f'<rect x="{x}" y="{y}" width="{pill_w}" height="72" rx="36" fill="white" fill-opacity="0.92"/>')
        parts.append(f'<circle cx="{x + 37}" cy="{y + 36}" r="19" fill="{rgb(ACCENTS[i % len(ACCENTS)])}"/>')
        parts.append(svg_text(x + 30, y + 44, [str(i + 1)], 20, "white", 700))
        parts.append(svg_text(x + 132, y + 46, [step], 29, "#1C222C", 700))
        if i < len(steps) - 1:
            parts.append(f'<line x1="{x + pill_w + 8}" y1="{y + 36}" x2="{x + pill_w + pill_gap - 24}" y2="{y + 36}" stroke="#97A4B5" stroke-width="4"/>')
            parts.append(f'<polygon points="{x + pill_w + pill_gap - 10},{y + 36} {x + pill_w + pill_gap - 24},{y + 29} {x + pill_w + pill_gap - 24},{y + 43}" fill="#97A4B5"/>')
    parts.append(svg_text(150, H - 58, ["架构提炼自《Ai 学习助手项目-功能确认单 V4》"], 24, "#748090"))
    parts.append(svg_text(W - 520, H - 58, ["PPT 立项版 · 16:9 高清图"], 24, "#748090"))
    parts.append("</svg>")
    SVG.write_text("\n".join(parts), encoding="utf-8")


def write_philosophy():
    philosophy = """# 清晰中枢

清晰中枢把复杂系统压缩为可被一眼理解的秩序。它不追求堆叠信息，而是让路径、边界与责任自然显形：入口在左，能力在中，基础设施在右，所有模块都像经过长时间推敲后的精密构件，留白承担解释，线条承担叙事。

空间以横向推进为主，强调从用户动作到模型调用再到计量反馈的闭环。每个层级都必须被严格安放，尺寸、间距、连接线和标签像工程图一样克制准确，呈现出被反复校准后的专业感，而不是临时拼装的汇报素材。

色彩使用低饱和底色与少量高识别度功能色。颜色不是装饰，而是信息分类系统：用户、接入、网关、服务、数据各有性格，但整体仍保持安静、可信、可交付。最终作品应当像由资深信息设计师精心打磨，细节有耐心，边界有判断。

文字只保留架构判断所必需的短标签。信息通过层级、卡片、箭头和闭环带表达，避免长段说明抢占视觉。观者无需阅读需求文档，也能理解系统如何从页面体验连接到模型平台、数据资产与安全运维。

所有形状都应当有明确功能：面板承载责任边界，卡片承载能力单元，箭头承载调用方向，底部闭环承载商业逻辑。最终画面必须看起来像经历过多轮细修，足够简洁，也足够立项场合使用。
"""
    MD.write_text(philosophy, encoding="utf-8")


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    draw_png()
    draw_svg()
    write_philosophy()
    print(PNG)
    print(SVG)
    print(MD)


if __name__ == "__main__":
    main()
